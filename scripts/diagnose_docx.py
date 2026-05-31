"""对比模板与输出文件的格式差异。"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE


def analyze(path: str) -> None:
    p = Path(path)
    print("=" * 70)
    print("文件:", p)
    if not p.exists():
        print("不存在")
        return

    doc = Document(str(p))
    print(f"段落数: {len(doc.paragraphs)}")
    print(f"节数: {len(doc.sections)}")

    if doc.sections:
        s = doc.sections[0]
        print(
            "页边距(cm): 上={:.2f} 下={:.2f} 左={:.2f} 右={:.2f}".format(
                s.top_margin.cm,
                s.bottom_margin.cm,
                s.left_margin.cm,
                s.right_margin.cm,
            )
        )

    para_styles = sorted(
        st.name for st in doc.styles if st.type == WD_STYLE_TYPE.PARAGRAPH and st.name
    )
    print(f"段落样式({len(para_styles)}个):")
    for name in para_styles[:30]:
        st = doc.styles[name]
        base = st.base_style.name if st.base_style else None
        print(f"  - {name!r} (based on: {base})")
    if len(para_styles) > 30:
        print(f"  ... 还有 {len(para_styles) - 30} 个")

    print("\n前 20 个非空段落:")
    count = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        count += 1
        if count > 20:
            break
        style_name = para.style.name if para.style else "?"
        pf = para.paragraph_format
        line_sp = pf.line_spacing
        first_indent = pf.first_line_indent.cm if pf.first_line_indent else 0
        run = para.runs[0] if para.runs else None
        font_name = run.font.name if run else None
        font_size = run.font.size.pt if run and run.font.size else None
        print(f"  {count}. 样式={style_name!r} | 字体={font_name} 字号={font_size}")
        print(f"      行距={line_sp} 首行缩进={first_indent:.2f}cm")
        print(f"      文本: {text[:60]}")


if __name__ == "__main__":
    template = sys.argv[1] if len(sys.argv) > 1 else r"templates\111.docx"
    output = sys.argv[2] if len(sys.argv) > 2 else r"output\from_template.docx"
    analyze(template)
    print()
    analyze(output)
