from docx import Document


def create_sample(path: str) -> None:
    doc = Document()
    doc.add_paragraph("论文标题示例")
    doc.add_paragraph("摘要")
    doc.add_paragraph(
        "这是摘要正文。本文研究了一种基于规则识别的论文排版方法，可用于快速统一 Word 文档格式。"
    )
    doc.add_paragraph("关键词：论文排版；Word；自动化")
    doc.add_paragraph("第1章 绪论")
    doc.add_paragraph(
        "绪论段落示例。许多学校对字体、行距、页边距有明确要求，手工调整非常耗时。"
    )
    doc.add_paragraph("1.1 研究背景")
    doc.add_paragraph(
        "研究背景段落示例。通过配置文件可以模拟不同学校的排版规范。"
    )
    doc.add_paragraph("1.1.1 问题提出")
    doc.add_paragraph("三级标题下的正文段落示例。")
    doc.add_paragraph("第2章 系统设计")
    doc.add_paragraph("第二章正文段落示例。")
    doc.save(path)


if __name__ == "__main__":
    create_sample("examples/sample_input.docx")
    print("已生成 examples/sample_input.docx")
