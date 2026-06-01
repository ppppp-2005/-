from __future__ import annotations

import re
from copy import deepcopy

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _para_is_centered(para: Paragraph) -> bool:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    align = para.alignment
    if align is None:
        align = para.paragraph_format.alignment
    return align == WD_ALIGN_PARAGRAPH.CENTER


def find_exemplar_paragraphs(template_doc, start_index: int = 0) -> dict[str, Paragraph]:
    """在模板【正文区】里找「范例段落」，用于复制真实格式（含手工设置的字体）。"""
    exemplars: dict[str, Paragraph] = {}

    for para in template_doc.paragraphs[start_index:]:
        text = para.text.strip()
        if not text:
            continue

        if "heading1" not in exemplars and re.search(
            r"^第\s*[\d一二三四五六七八九十百]+\s*章\s+\S", text
        ):
            exemplars["heading1"] = para
        elif "heading2" not in exemplars and re.match(r"^\d+\.\d+\s+\S", text):
            exemplars["heading2"] = para
        elif "heading3" not in exemplars and re.match(r"^\d+\.\d+\.\d+\s+\S", text):
            exemplars["heading3"] = para
        elif "abstract_title" not in exemplars and re.match(
            r"^摘\s*要$|^Abstract$", text, re.I
        ):
            exemplars["abstract_title"] = para
        elif "keywords" not in exemplars and re.match(
            r"^关键词[：:]|^Keywords[：:]", text, re.I
        ):
            exemplars["keywords"] = para
        elif "body" not in exemplars and len(text) >= 15 and not _para_is_centered(para):
            exemplars["body"] = para

    return exemplars


def _copy_east_asia_font(source_run, target_run) -> None:
    src_rpr = source_run._element.rPr
    if src_rpr is None:
        return
    rfonts = src_rpr.rFonts
    if rfonts is None:
        return
    east = rfonts.get(qn("w:eastAsia"))
    if east:
        tgt_rpr = target_run._element.get_or_add_rPr()
        tgt_rfonts = tgt_rpr.get_or_add_rFonts()
        tgt_rfonts.set(qn("w:eastAsia"), east)


def clone_run_format(source_run, target_run) -> None:
    target_run.font.name = source_run.font.name
    target_run.font.size = source_run.font.size
    target_run.font.bold = source_run.font.bold
    target_run.font.italic = source_run.font.italic
    target_run.font.underline = source_run.font.underline
    try:
        if source_run.font.color and source_run.font.color.rgb:
            target_run.font.color.rgb = source_run.font.color.rgb
    except (AttributeError, TypeError):
        pass
    _copy_east_asia_font(source_run, target_run)

    src_rpr = source_run._element.rPr
    if src_rpr is not None:
        tgt_rpr = target_run._element.get_or_add_rPr()
        for child in src_rpr:
            tag = child.tag.split("}")[-1]
            if tag in ("rFonts", "sz", "szCs", "b", "i", "u", "color", "highlight"):
                existing = tgt_rpr.find(qn(f"w:{tag}"))
                if existing is not None:
                    tgt_rpr.remove(existing)
                tgt_rpr.append(deepcopy(child))


LABEL_PREFIX_RE = re.compile(
    r"^(摘\s*要\s*[：:]|Abstract\s*[：:]|关键词\s*[：:]|关键字\s*[：:]|Keywords\s*[：:]|Key\s+words\s*[：:])",
    re.I,
)


def _split_label_text(text: str) -> tuple[str, str] | None:
    match = LABEL_PREFIX_RE.match(text.strip())
    if not match:
        return None
    label = match.group(1)
    body = text.strip()[match.end() :].lstrip()
    return label, body


def _first_content_run_after(source: Paragraph, start_index: int):
    for run in source.runs[start_index + 1 :]:
        if run.text.strip():
            return run
    return None


def _clone_labelled_run_format(source: Paragraph, target: Paragraph) -> bool:
    target_parts = _split_label_text(target.text)
    if not target_parts or len(source.runs) < 2:
        return False

    label_run = None
    label_run_index = -1
    for index, run in enumerate(source.runs):
        if _split_label_text(run.text):
            label_run = run
            label_run_index = index
            break
    if label_run is None:
        return False

    body_run = _first_content_run_after(source, label_run_index)
    if body_run is None:
        return False

    label, body = target_parts
    target.clear()
    new_label = target.add_run(label)
    new_body = target.add_run(body)
    clone_run_format(label_run, new_label)
    clone_run_format(body_run, new_body)
    return True


def clone_paragraph_format(source: Paragraph, target: Paragraph) -> None:
    """完整复制段落格式：样式 + 段落属性 + 字体（含中文字体）。"""
    target.style = source.style

    sp = source.paragraph_format
    tp = target.paragraph_format
    tp.alignment = sp.alignment
    tp.left_indent = sp.left_indent
    tp.right_indent = sp.right_indent
    tp.first_line_indent = sp.first_line_indent
    tp.line_spacing = sp.line_spacing
    tp.line_spacing_rule = sp.line_spacing_rule
    tp.space_before = sp.space_before
    tp.space_after = sp.space_after
    tp.keep_together = sp.keep_together
    tp.keep_with_next = sp.keep_with_next
    tp.widow_control = sp.widow_control
    tp.page_break_before = sp.page_break_before

    # 复制段落级 XML（缩进、行距等更完整）
    if source._element.pPr is not None:
        new_ppr = deepcopy(source._element.pPr)
        old = target._element.pPr
        if old is not None:
            target._element.remove(old)
        target._element.insert(0, new_ppr)

    if not source.runs or not target.runs:
        return

    if _clone_labelled_run_format(source, target):
        return

    src_run = next((run for run in source.runs if run.text.strip()), source.runs[0])
    for run in target.runs:
        clone_run_format(src_run, run)


def exemplar_style_label(para: Paragraph) -> str:
    text = para.text.strip()
    preview = text[:20] + ("..." if len(text) > 20 else "")
    style = para.style.name if para.style and para.style.name else "?"
    return f"{style}「{preview}」"
