from __future__ import annotations

import re
from collections import Counter

from docx import Document

from .classifier import classify_paragraph
from .style_constants import STYLE_CANDIDATES


def styles_used_in_template(doc: Document) -> set[str]:
    return {
        para.style.name
        for para in doc.paragraphs
        if para.text.strip() and para.style and para.style.name
    }


def build_template_style_profile(
    doc: Document, start_index: int = 0
) -> dict[str, str | None]:
    """从模板【正文区】实际出现的段落，推断各级标题/正文对应的 Word 样式名。"""
    used = styles_used_in_template(doc)
    if not used:
        return {}

    counters: dict[str, Counter[str]] = {
        "body": Counter(),
        "heading1": Counter(),
        "heading2": Counter(),
        "heading3": Counter(),
        "heading4": Counter(),
        "abstract_title": Counter(),
        "abstract_body": Counter(),
        "keywords": Counter(),
        "references_title": Counter(),
        "references_body": Counter(),
        "acknowledgment": Counter(),
        "figure_caption": Counter(),
        "table_caption": Counter(),
    }

    in_abstract = False
    in_references = False

    for para in doc.paragraphs[start_index:]:
        text = para.text.strip()
        if not text or not para.style or not para.style.name:
            continue
        name = para.style.name

        if re.match(r"^摘\s*要$|^Abstract$", text, re.I):
            counters["abstract_title"][name] += 1
            in_abstract = True
            in_references = False
            continue
        if re.match(r"^关键词[：:]|^Keywords[：:]", text, re.I):
            counters["keywords"][name] += 1
            continue
        if re.match(r"^参考文献$|^References$", text, re.I):
            counters["references_title"][name] += 1
            in_references = True
            in_abstract = False
            continue
        if re.match(r"^致\s*谢$", text, re.I):
            counters["acknowledgment"][name] += 1
            continue
        if re.match(r"^图[\d\-\.]+", text) or re.match(r"^Figure\s+\d+", text, re.I):
            counters["figure_caption"][name] += 1
            continue
        if re.match(r"^表[\d\-\.]+", text) or re.match(r"^Table\s+\d+", text, re.I):
            counters["table_caption"][name] += 1
            continue
        if in_abstract and len(text) >= 20 and not re.match(
            r"^关键词[：:]|^Keywords[：:]", text, re.I
        ):
            counters["abstract_body"][name] += 1
        if in_references and re.match(r"^\[\d+\]", text):
            counters["references_body"][name] += 1
        if re.search(r"^第\s*[\d一二三四五六七八九十百]+\s*章\s+\S", text):
            counters["heading1"][name] += 1
        elif re.match(r"^\d+\.\d+\.\d+\.\d+\s+\S", text):
            counters["heading4"][name] += 1
        elif re.match(r"^\d+\.\d+\.\d+\s+\S", text):
            counters["heading3"][name] += 1
        elif re.match(r"^\d+\.\d+\s+\S", text):
            counters["heading2"][name] += 1
        elif len(text) >= 20:
            counters["body"][name] += 1

        classify_paragraph(text, [])

    def pick(counter: Counter[str], fallback_keywords: tuple[str, ...]) -> str | None:
        if counter:
            return counter.most_common(1)[0][0]
        for style_name in used:
            lower = style_name.lower()
            if any(k in lower or k in style_name for k in fallback_keywords):
                return style_name
        return None

    profile = {
        "body": pick(counters["body"], ("正文", "normal", "body")),
        "heading1": pick(counters["heading1"], ("标题1", "标题 1", "heading 1", "一级标题", "1级标题")),
        "heading2": pick(counters["heading2"], ("标题2", "标题 2", "heading 2", "二级标题", "2级标题")),
        "heading3": pick(counters["heading3"], ("标题3", "标题 3", "heading 3", "三级标题", "3级标题")),
        "heading4": pick(counters["heading4"], ("标题4", "标题 4", "heading 4", "四级标题", "4级标题")),
        "abstract_title": pick(counters["abstract_title"], ("摘要", "abstract")),
        "abstract_body": pick(counters["abstract_body"], ("摘要正文", "abstract body")),
        "keywords": pick(counters["keywords"], ("关键词", "keyword")),
        "references_title": pick(counters["references_title"], ("参考文献", "references")),
        "references_body": pick(counters["references_body"], ("参考文献", "reference")),
        "acknowledgment": pick(counters["acknowledgment"], ("致谢", "acknowledgment")),
        "figure_caption": pick(counters["figure_caption"], ("图标题", "caption")),
        "table_caption": pick(counters["table_caption"], ("表标题", "table caption")),
    }

    if not profile["body"]:
        all_used = Counter(
            para.style.name for para in doc.paragraphs if para.text.strip() and para.style
        )
        if all_used:
            profile["body"] = all_used.most_common(1)[0][0]

    return profile


def merge_profile_with_exemplars(
    profile: dict[str, str | None], format_library: dict
) -> dict[str, str | None]:
    """把格式范例段落里的样式名补进 profile，避免模板只有样式定义、没有正文范例时漏套。"""
    merged = dict(profile)
    for para_type, exemplar in format_library.items():
        if merged.get(para_type):
            continue
        style_name = exemplar.style.name if exemplar.style else None
        if style_name:
            merged[para_type] = style_name
    return merged


def resolve_style_with_profile(
    para_type: str,
    profile: dict[str, str | None],
    available: set[str],
    used_in_template: set[str],
) -> str | None:
    """优先用模板里真实用过的样式，避免误选 Word 内置的 Empty Heading 1。"""
    if para_type in profile and profile[para_type] and profile[para_type] in available:
        return profile[para_type]

    for candidate in STYLE_CANDIDATES.get(para_type, STYLE_CANDIDATES["body"]):
        if candidate in used_in_template and candidate in available:
            return candidate

    for candidate in STYLE_CANDIDATES.get(para_type, STYLE_CANDIDATES["body"]):
        if candidate in available:
            return candidate

    if profile.get("body") in available:
        return profile["body"]
    if "Normal" in available:
        return "Normal"
    return next(iter(available), None)
